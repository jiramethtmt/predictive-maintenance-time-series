import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "predictive-maintenance-executive.pptx"

REPO = "https://github.com/jiramethtmt/predictive-maintenance-time-series"
NOTEBOOK_PATH = "notebooks/01_eda_scania_component_x.ipynb"
NOTEBOOK_VIEWER = f"https://nbviewer.org/github/jiramethtmt/predictive-maintenance-time-series/blob/main/{NOTEBOOK_PATH}"
NOTEBOOK_SOURCE = f"{REPO}/blob/main/{NOTEBOOK_PATH}"
DATASET_DOI = "https://doi.org/10.5878/jvb5-d390"

INK = RGBColor(0x0B, 0x25, 0x45)
BODY = RGBColor(0x34, 0x3D, 0x4A)
MUTED = RGBColor(0x8A, 0x94, 0xA0)
ACCENT = RGBColor(0x1F, 0x7A, 0x6B)
WARN = RGBColor(0xB5, 0x48, 0x1E)
RULE = RGBColor(0xDD, 0xE3, 0xE9)
TINT = RGBColor(0xF4, 0xF7, 0xF9)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.9)
CONTENT_TOP = Inches(2.25)
COLUMN = W - 2 * MARGIN


def textframe(slide, left, top, width, height):
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return frame


def line(frame, text, size, color, bold=False, after=6, align=PP_ALIGN.LEFT, space=1.0):
    blank = len(frame.paragraphs) == 1 and not frame.paragraphs[0].runs
    para = frame.paragraphs[0] if blank else frame.add_paragraph()
    para.alignment = align
    para.space_after = Pt(after)
    para.line_spacing = space
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return para


def rule(slide, top, width=None, color=RULE, thickness=Pt(1.25)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top, width or COLUMN, thickness)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def slide_of(deck, kicker, title, number):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    head = textframe(slide, MARGIN, Inches(0.6), COLUMN, Inches(1.2))
    line(head, kicker.upper(), 12, ACCENT, bold=True, after=7)
    line(head, title, 31, INK, bold=True, after=0)
    rule(slide, Inches(1.95))
    foot = textframe(slide, MARGIN, H - Inches(0.62), COLUMN, Inches(0.3))
    line(foot, f"SCANIA Component X   |   {number:02d}", 10, MUTED, after=0)
    return slide


def bullets(slide, top, items, size=16, gap=14):
    frame = textframe(slide, MARGIN, top, COLUMN, H - top - Inches(1.0))
    for lead, rest in items:
        para = line(frame, lead, size, INK, bold=True, after=2, space=1.15)
        if rest:
            run = para.add_run()
            run.text = "  " + rest
            run.font.size = Pt(size)
            run.font.color.rgb = BODY
            run.font.name = FONT
            para.space_after = Pt(gap)
    return frame


def stats(slide, top, cards):
    gap = Inches(0.35)
    width = int((COLUMN - gap * (len(cards) - 1)) / len(cards))
    for index, (value, label, color) in enumerate(cards):
        left = MARGIN + index * (width + gap)
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(1.65))
        panel.fill.solid()
        panel.fill.fore_color.rgb = TINT
        panel.line.fill.background()
        panel.shadow.inherit = False
        panel.adjustments[0] = 0.05
        frame = textframe(slide, left + Inches(0.3), top + Inches(0.28), width - Inches(0.6), Inches(1.2))
        line(frame, value, 34, color, bold=True, after=4)
        line(frame, label, 12, MUTED, after=0, space=1.1)


def table(slide, top, header, rows, widths, highlight=0):
    height = Inches(0.46) * (len(rows) + 1)
    grid = slide.shapes.add_table(len(rows) + 1, len(header), MARGIN, top, COLUMN, height).table
    for index, share in enumerate(widths):
        grid.columns[index].width = int(COLUMN * share)
    for index, cells in enumerate([header] + rows):
        grid.rows[index].height = Inches(0.46)
        strong = index == highlight
        for column, text in enumerate(cells):
            cell = grid.cell(index, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TINT if index == 0 else PAPER
            cell.margin_left = cell.margin_right = Inches(0.14)
            frame = cell.text_frame
            frame.word_wrap = True
            line(
                frame,
                text,
                13 if index == 0 else 15,
                MUTED if index == 0 else (ACCENT if strong else BODY),
                bold=index == 0 or strong,
                after=0,
                align=PP_ALIGN.RIGHT if column else PP_ALIGN.LEFT,
            )


def links(slide, top, entries):
    for index, (label, caption, url) in enumerate(entries):
        offset = top + index * Inches(1.15)
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, offset, COLUMN, Inches(0.95))
        panel.fill.solid()
        panel.fill.fore_color.rgb = TINT
        panel.line.fill.background()
        panel.shadow.inherit = False
        panel.adjustments[0] = 0.08
        frame = textframe(slide, MARGIN + Inches(0.35), offset + Inches(0.16), COLUMN - Inches(0.7), Inches(0.7))
        para = line(frame, label, 16, INK, bold=True, after=2)
        para.runs[0].hyperlink.address = url
        line(frame, caption, 12, MUTED, after=0)


def title_slide(deck):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), H)
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    band.shadow.inherit = False
    frame = textframe(slide, Inches(1.1), Inches(2.3), Inches(10.6), Inches(3.0))
    line(frame, "PREDICTIVE MAINTENANCE", 13, ACCENT, bold=True, after=14)
    line(frame, "Calling trucks in before the", 40, INK, bold=True, after=2, space=1.05)
    line(frame, "component fails", 40, INK, bold=True, after=20, space=1.05)
    line(
        frame,
        "A cost-optimal failure-risk model on the SCANIA Component X fleet, and what it would take to run it.",
        16,
        BODY,
        after=0,
        space=1.25,
    )
    rule(slide, Inches(5.5), width=Inches(2.2), color=RULE)
    tail = textframe(slide, Inches(1.1), Inches(5.8), Inches(10.6), Inches(0.8))
    line(tail, "Executive summary", 13, MUTED, after=3)
    line(tail, "33,000 trucks   |   IDA 2024 challenge cost matrix   |   Three-seed result", 13, MUTED, after=0)
    return slide


def build(out=OUT):
    deck = Presentation()
    deck.slide_width, deck.slide_height = W, H

    title_slide(deck)

    slide = slide_of(deck, "The question", "What the model is actually asked", 2)
    bullets(
        slide,
        CONTENT_TOP,
        [
            ("For one truck, right now:", "how close is the monitored component to failure, using only the readouts that truck has already sent home."),
            ("The output is an action, not a score:", "call this truck into the workshop, or leave it running."),
            ("Five degradation classes:", "class 0 is more than 48 time steps from failure, class 4 is 0 to 6 steps away."),
            ("Only 2.7% of the fleet is at risk", "in any scoring window. The rest are healthy and must not be disturbed."),
        ],
    )

    slide = slide_of(deck, "Why accuracy fails", "The economics decide the model, not the metric", 3)
    stats(
        slide,
        CONTENT_TOP,
        [
            ("500", "Cost of missing a truck that fails", WARN),
            ("10", "Cost of a workshop check that finds nothing", ACCENT),
            ("50 : 1", "False alarms one real catch pays for", INK),
        ],
    )
    bullets(
        slide,
        CONTENT_TOP + Inches(2.1),
        [
            ("A model scored on accuracy predicts healthy for everyone,", "is right 97% of the time, and saves nothing."),
            ("We minimise expected cost instead.", "The model emits probabilities; the decision rule picks the cheapest action under the published cost matrix."),
            ("Low precision is correct here, not broken.", "Optimising precision on this loss destroys value."),
        ],
    )

    slide = slide_of(deck, "Result", "Cost against every alternative policy", 4)
    stats(
        slide,
        CONTENT_TOP,
        [
            ("35,469 \u00b1 849", "Test cost, mean of three seeds", ACCENT),
            ("29%", "Cheaper than checking every truck", INK),
            ("0.77", "Share of at-risk trucks caught", INK),
        ],
    )
    table(
        slide,
        CONTENT_TOP + Inches(2.05),
        ["Policy", "Test cost"],
        [
            ["AutoGluon LightGBM, cost-optimal rule (3 seeds)", "35,469 \u00b1 849"],
            ["Check every truck", "49,671"],
            ["Act on the most likely class", "56,100"],
            ["Check nothing", "56,100"],
        ],
        [0.72, 0.28],
        highlight=1,
    )

    slide = slide_of(deck, "In operational terms", "What the number buys the workshop", 5)
    bullets(
        slide,
        CONTENT_TOP,
        [
            ("109 of 142 at-risk trucks are caught", "before the component fails, at the tuned operating point."),
            ("About 2,100 of 5,045 trucks are called in.", "Checking all 5,045 catches every failure and costs 40% more."),
            ("The saving comes from not inspecting the other 2,900,", "while still reaching most of the trucks that were going to fail."),
            ("33 at-risk trucks are still missed.", "That residual is the honest cost of the current model, and it is where further work pays."),
        ],
    )

    slide = slide_of(deck, "Benchmark", "Every published result on this dataset", 6)
    table(
        slide,
        CONTENT_TOP,
        ["Approach", "Test cost"],
        [
            ["Our work, AutoGluon LightGBM (3 seeds)", "35,469 \u00b1 849"],
            ["CatBoost, same study, best row read off the test set", "36,724"],
            ["XGBoost, empirical study, the model they actually selected", "37,733"],
            ["Bi-LSTM, Zhong and Wang, IDA 2024", "39,123"],
            ["GNN, Parton et al., IDA 2024", "47,612"],
            ["XGBoost, Carpentier et al., IDA 2024, level with checking every truck", "49,671"],
        ],
        [0.72, 0.28],
        highlight=1,
    )
    bullets(
        slide,
        CONTENT_TOP + Inches(3.35),
        [
            ("The honest target is 37,733, the model the leading study selected on validation.", "We are 2,264 ahead of it, and that margin is wider than our own seed spread."),
            ("No published entry states how it converts probabilities into a decision,", "and on our own model that choice is worth 20,631, more than the whole spread of this table."),
        ],
        size=14,
        gap=8,
    )
    slide = slide_of(deck, "Confidence", "Why this number is reportable", 7)
    bullets(
        slide,
        CONTENT_TOP,
        [
            ("Three seeds, not one.", "A single run on this data moves by more than most of the improvements teams chase on it."),
            ("The noise floor was measured, not assumed.", "Single-split tuning had a seed-to-seed spread of 1,668. Pooling validation into training and tuning on out-of-fold predictions halved it to 849."),
            ("A leakage audit was run and it found something.", "A backward fill let a scoring point read readouts from its own future. Removing it moved the result from a flattering 35,925 to an honest 39,379, before the protocol work earned the gain back."),
            ("The test set was scored once,", "at the end, after every decision was fixed on training data."),
        ],
        size=15,
        gap=13,
    )

    slide = slide_of(deck, "Limitation", "Ranking is the binding constraint", 8)
    stats(
        slide,
        CONTENT_TOP,
        [
            ("15 / 142", "At-risk trucks inside the top 100 by risk", WARN),
            ("47 / 142", "Inside the top 500", INK),
            ("2,100", "Alarms the cost rule actually raises", INK),
        ],
    )
    bullets(
        slide,
        CONTENT_TOP + Inches(2.1),
        [
            ("The cost win comes from alarming broadly, not from sharp ranking.", ""),
            ("If the workshop can only take 100 trucks a day,", "this model finds roughly one in ten of the failures, not eight in ten."),
            ("Capacity-constrained rollout needs better ordering,", "which is the next piece of modelling work, not more threshold tuning."),
        ],
    )

    slide = slide_of(deck, "Running it for real", "The loop this would sit inside", 9)
    bullets(
        slide,
        CONTENT_TOP,
        [
            ("Batch scoring, not a real-time alarm.", "These are aggregate counters uploaded per cycle, not a live sensor stream. Trucks are rescored when new readouts arrive."),
            ("Output is a ranked work list for the planner,", "trimmed to what the workshop can absorb that day."),
            ("Every dispatched truck returns a label.", "Found or not found feeds the next training round; without that loop the model stops improving on day one."),
            ("Drift has to be watched.", "New vehicle configurations and new routes shift the inputs; retraining is scheduled, not incidental."),
            ("Each alert needs a reason a mechanic will accept,", "such as the counter that accelerated threefold over three cycles."),
        ],
        size=15,
        gap=12,
    )

    slide = slide_of(deck, "Method, in one slide", "What kind of model this is", 10)
    bullets(
        slide,
        CONTENT_TOP,
        [
            ("Discrete-time survival classification on panel data.", "Not a sequence forecaster: the history up to a scoring point is compressed into features, and a gradient-boosted tree predicts the degradation class."),
            ("Features are rates, not levels.", "The counters are cumulative and never reset, so wear speed over 3, 10 and 20 readout windows carries the signal, plus acceleration, histogram drift and reporting gaps."),
            ("89% of outcomes are censored.", "Most trucks never failed inside the observation window, which is why a survival framing beats a plain binary classifier."),
            ("Folds are grouped by vehicle,", "so near-duplicate scoring points from one truck can never sit on both sides of a split."),
        ],
        size=15,
        gap=13,
    )

    slide = slide_of(deck, "Go deeper", "The analysis behind these numbers", 11)
    links(
        slide,
        CONTENT_TOP,
        [
            ("Open the full EDA notebook", NOTEBOOK_VIEWER, NOTEBOOK_VIEWER),
            ("Repository, code and reproduction steps", REPO, REPO),
            ("SCANIA Component X dataset, CC BY 4.0", DATASET_DOI, DATASET_DOI),
            ("Notebook source on GitHub", NOTEBOOK_SOURCE, NOTEBOOK_SOURCE),
        ],
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(out)
    return out


if __name__ == "__main__":
    print(build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT))